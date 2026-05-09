"""Core conversion logic. Each function takes input_path -> output_path."""
import subprocess, os, shutil
from pathlib import Path
from src.app.config import settings
from src.app.logging import logger


def _libreoffice_convert(input_path: str, output_dir: str, output_format: str) -> str:
    """Use LibreOffice headless to convert docs. Returns output file path."""
    cmd = [
        settings.LIBREOFFICE_PATH, "--headless", "--convert-to",
        output_format, "--outdir", output_dir, input_path,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice failed: {result.stderr}")
    stem = Path(input_path).stem
    return os.path.join(output_dir, f"{stem}.{output_format}")


def convert(job_id: str, input_path: str, input_fmt: str, output_fmt: str) -> str:
    """Main dispatch function. Returns output filename (not full path)."""
    output_dir = settings.OUTPUT_DIR
    stem = job_id
    output_filename = f"{stem}.{output_fmt}"
    output_path = os.path.join(output_dir, output_filename)

    logger.info(f"Converting [{input_fmt}] -> [{output_fmt}] | job={job_id}")

    # ── LibreOffice-backed conversions ─────────────────────────────────
    lo_pairs = {
        ("docx", "pdf"), ("docx", "html"), ("docx", "txt"),
        ("pptx", "pdf"), ("ppt",  "pdf"),
        ("xlsx", "pdf"), ("xls",  "pdf"),
        ("txt",  "pdf"), ("html", "pdf"),
        ("odt",  "pdf"), ("odp",  "pdf"),
    }
    if (input_fmt, output_fmt) in lo_pairs:
        result_path = _libreoffice_convert(input_path, output_dir, output_fmt)
        final = os.path.join(output_dir, output_filename)
        if result_path != final:
            shutil.move(result_path, final)
        return output_filename

    # ── PDF -> DOCX via pdf2docx ────────────────────────────────────────
    if input_fmt == "pdf" and output_fmt == "docx":
        from pdf2docx import Converter as P2D
        cv = P2D(input_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
        return output_filename

    # ── PDF -> TXT via pdfminer ─────────────────────────────────────────
    if input_fmt == "pdf" and output_fmt == "txt":
        from pdfminer.high_level import extract_text
        text = extract_text(input_path)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)
        return output_filename

    # ── PDF -> Images via pdf2image ─────────────────────────────────────
    if input_fmt == "pdf" and output_fmt in ("png", "jpg"):
        from pdf2image import convert_from_path
        fmt_map = {"png": "PNG", "jpg": "JPEG"}
        images = convert_from_path(input_path, dpi=150)
        if len(images) == 1:
            images[0].save(output_path, fmt_map[output_fmt])
        else:
            import zipfile
            zip_path = os.path.join(output_dir, f"{stem}.zip")
            with zipfile.ZipFile(zip_path, "w") as zf:
                for i, img in enumerate(images):
                    tmp = os.path.join(output_dir, f"{stem}_p{i+1}.{output_fmt}")
                    img.save(tmp, fmt_map[output_fmt])
                    zf.write(tmp, arcname=os.path.basename(tmp))
                    os.remove(tmp)
            return f"{stem}.zip"
        return output_filename

    # ── Image -> PDF via Pillow ─────────────────────────────────────────
    if input_fmt in ("png", "jpg", "jpeg", "bmp") and output_fmt == "pdf":
        from PIL import Image
        img = Image.open(input_path).convert("RGB")
        img.save(output_path, "PDF")
        return output_filename

    # ── XLSX -> CSV via openpyxl ────────────────────────────────────────
    if input_fmt in ("xlsx", "xls") and output_fmt == "csv":
        import openpyxl, csv
        wb = openpyxl.load_workbook(input_path, read_only=True)
        ws = wb.active
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(row)
        return output_filename

    # ── CSV -> XLSX via openpyxl ────────────────────────────────────────
    if input_fmt == "csv" and output_fmt == "xlsx":
        import openpyxl, csv
        wb = openpyxl.Workbook()
        ws = wb.active
        with open(input_path, newline="", encoding="utf-8") as f:
            for row in csv.reader(f):
                ws.append(row)
        wb.save(output_path)
        return output_filename

    # ── TXT -> DOCX via python-docx ─────────────────────────────────────
    if input_fmt == "txt" and output_fmt == "docx":
        from docx import Document
        doc = Document()
        with open(input_path, encoding="utf-8") as f:
            for line in f:
                doc.add_paragraph(line.rstrip())
        doc.save(output_path)
        return output_filename

    raise ValueError(f"No converter found for {input_fmt} -> {output_fmt}")
